"""2.7.61: firma en el Mostrador (6), quién tiene qué (7), resumen semanal (5) e
importar formaciones y reconocimientos desde Excel (8)."""
import io
from datetime import date, datetime

import main
from auth import hash_password
from models import (AlbaranSalida, Almacen, Aviso, FormacionTrabajador, Herramienta, Material, ReconocimientoMedico,
                    Trabajador, Usuario)
from security import generar_csrf_token

PNG_1PX = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="


def _setup(db):
    almacen = Almacen(nombre="Nave m58", codigo="MRD-M58", activo=True)
    db.add(almacen)
    db.flush()
    db.add(Usuario(username="admin-m58", password_hash=hash_password("ClaveSegura123!"), nombre="Admin", rol="admin", activo=True, must_change_password=False, almacen_id=almacen.id))
    t = Trabajador(nombre="Erick", apellidos="Cinco", activo=True, almacen_id=almacen.id, codigo="T-058", telefono="600123456")
    db.add(t)
    db.flush()
    h1 = Herramienta(codigo="M58-H1", nombre="Taladro cinco", estado="disponible", activa=True, almacen_id=almacen.id, precio_compra=250)
    h2 = Herramienta(codigo="M58-H2", nombre="Radial cinco", estado="disponible", activa=True, almacen_id=almacen.id, precio_compra=120)
    db.add_all([h1, h2, Material(codigo="M58-M1", nombre="Discos cinco", unidad="ud", stock_actual=1, stock_minimo=5, activo=True, almacen_id=almacen.id)])
    db.commit()
    return almacen, t, h1, h2


def _login(client):
    resp = client.post("/login", data={"username": "admin-m58", "password": "ClaveSegura123!"}, follow_redirects=False)
    client.cookies.set("mrd_token", resp.cookies["mrd_token"])
    token = generar_csrf_token()
    client.cookies.set("mrd_csrf", token)
    return {"X-CSRF-Token": token, "Accept": "application/json"}


def _salida(client, h, almacen, t, ids, **extra):
    body = {"operacion_id": f"m58-op-{ids[0]}", "accion": "salida", "trabajador_id": t.id, "almacen_id": almacen.id,
            "lineas": [{"tipo": "herramienta", "id": i, "cantidad": 1} for i in ids]}
    body.update(extra)
    return client.post("/api/mostrador/operar", json=body, headers=h)


def test_firma_en_mostrador_queda_en_el_albaran(client, db):
    almacen, t, h1, h2 = _setup(db)
    h = _login(client)
    r = _salida(client, h, almacen, t, [h1.id], firma_datos=PNG_1PX, firma_nombre="Erick Cinco")
    assert r.status_code == 200, r.text
    alb = db.get(AlbaranSalida, r.json()["albaran_id"])
    assert alb.firma_datos and alb.firma_datos.startswith("data:image/png") and alb.firma_nombre == "Erick Cinco"
    r = _salida(client, h, almacen, t, [h2.id], firma_datos="no-es-una-imagen")
    assert r.status_code == 400
    assert 'id="firma-canvas"' in client.get("/mostrador").text


def test_quien_tiene_que_y_excel(client, db):
    almacen, t, h1, h2 = _setup(db)
    h = _login(client)
    assert _salida(client, h, almacen, t, [h1.id, h2.id]).status_code == 200
    datos = main._quien_tiene_que(db, almacen.id)
    assert datos["total"] == 2 and datos["personas"] == 1 and datos["valor_total"] == 370
    g = datos["grupos"][0]
    assert g["nombre"] == "Erick Cinco" and g["n"] == 2 and g["items"][0]["desde"] == date.today().strftime("%d/%m/%Y")
    page = client.get("/informes/quien-tiene-que")
    assert page.status_code == 200 and "Erick Cinco" in page.text and "wa.me/34600123456" in page.text and "Taladro cinco" in page.text
    xl = client.get("/informes/quien-tiene-que/excel")
    assert xl.status_code == 200 and xl.content[:2] == b"PK"
    assert 'href="/informes/quien-tiene-que"' in client.get("/informes").text


def test_resumen_semanal_pagina_aviso_y_push(client, db):
    almacen, t, h1, h2 = _setup(db)
    h = _login(client)
    _salida(client, h, almacen, t, [h1.id])
    titulo, texto, datos = main._resumen_semanal_texto(db, almacen.id)
    assert "QUIÉN TIENE QUÉ: 1 herramientas fuera" in texto and "Erick Cinco: 1" in texto and "STOCK BAJO: 1 materiales" in texto and datos["bajos"] == 1
    page = client.get("/resumen-semanal")
    assert page.status_code == 200 and "wa.me/?text=" in page.text and "Erick Cinco" in page.text
    r = client.post("/api/resumen-semanal/push", headers=h)
    assert r.status_code == 200 and r.json()["ok"] and r.json()["repetido"] is False
    assert db.query(Aviso).filter(Aviso.titulo.like("Resumen semanal %")).count() == 1
    assert main._resumen_semanal_publicar(db)["repetido"] is True
    assert 'href="/resumen-semanal"' in client.get("/").text


def test_importar_formaciones_y_reconocimientos(client, db):
    almacen, t, h1, h2 = _setup(db)
    h = _login(client)
    from openpyxl import Workbook, load_workbook
    r = client.get("/trabajadores/importar-formaciones/plantilla")
    assert r.status_code == 200 and [ws.title for ws in load_workbook(io.BytesIO(r.content)).worksheets] == ["Formaciones", "Reconocimientos"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Formaciones"
    ws.append(main.FORMACIONES_CABECERAS)
    ws.append(["T-058", "", "PRL 20 h", "PRL", "FLC", "01/09/2026", "01/09/2029", "C-1"])
    ws.append(["", "erick cinco", "Trabajos en altura", "", "", datetime(2026, 5, 2), date(2028, 5, 2), ""])
    ws.append(["NO-EXISTE", "", "Curso X", "", "", "", "", ""])
    ws2 = wb.create_sheet("Reconocimientos")
    ws2.append(main.RECONOCIMIENTOS_CABECERAS)
    ws2.append(["T-058", "", "10/01/2026", "Apto", "10/01/2027", "Dr. X", "Centro Y"])
    ws2.append(["T-058", "", "10/01/2026", "Apto", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/trabajadores/importar-formaciones", files={"archivo": ("f.xlsx", io.BytesIO(buf.getvalue()), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["formaciones"] == 2 and d["reconocimientos"] == 1 and d["saltadas"] == 1 and len(d["errores"]) == 1 and "NO-EXISTE" in d["errores"][0]
    f = db.query(FormacionTrabajador).filter(FormacionTrabajador.trabajador_id == t.id).order_by(FormacionTrabajador.id).all()
    assert [x.nombre_curso for x in f] == ["PRL 20 h", "Trabajos en altura"] and f[0].fecha_caducidad == date(2029, 9, 1) and f[1].fecha_realizacion == date(2026, 5, 2)
    rc = db.query(ReconocimientoMedico).filter(ReconocimientoMedico.trabajador_id == t.id).one()
    assert rc.resultado == "apto" and rc.fecha_proxima == date(2027, 1, 10)
    assert client.get("/trabajadores/importar-formaciones").status_code == 200
    assert 'href="/trabajadores/importar-formaciones"' in client.get("/trabajadores").text
    ev = [e for e in main._calendario_eventos(db, almacen.id) if e["cat"] == "trabajador"]
    assert len(ev) == 3
