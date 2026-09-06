"""2.7.48: informes de gasto por herramienta y consumo por obra, con Excel."""
import io
from datetime import datetime

from openpyxl import load_workbook

from auth import hash_password
from models import Almacen, Herramienta, MantenimientoProgramado, Material, MovimientoMaterial, Obra, Reparacion, Usuario
from security import generar_csrf_token


def test_gasto_por_herramienta_y_consumo_por_obra(client, db):
    almacen = Almacen(nombre="Nave gasto", codigo="MRD-GAS", activo=True)
    db.add(almacen)
    db.flush()
    db.add(Usuario(username="admin-gas", password_hash=hash_password("ClaveSegura123!"), nombre="Admin",
                   rol="admin", activo=True, must_change_password=False, almacen_id=almacen.id))
    h1 = Herramienta(codigo="GAS-1", nombre="Martillo caro", estado="disponible", activa=True, almacen_id=almacen.id, precio_compra=400)
    h2 = Herramienta(codigo="GAS-2", nombre="Radial barata", estado="entregada", activa=True, almacen_id=almacen.id, precio_compra=100)
    obra = Obra(numero="2026-0099", nombre="Obra consumo", activa=True, almacen_id=almacen.id)
    db.add_all([h1, h2, obra])
    db.flush()
    h2.obra_id = obra.id
    mat = Material(codigo="MAT-GAS", nombre="Discos", unidad="ud", stock_actual=50, activo=True, almacen_id=almacen.id)
    db.add(mat)
    db.flush()
    db.add_all([
        Reparacion(numero="REP-GAS-1", herramienta_id=h1.id, estado="finalizada", coste_final=150),
        Reparacion(numero="REP-GAS-2", herramienta_id=h1.id, estado="en_reparacion", coste_estimado=100),
        MantenimientoProgramado(tipo_activo="herramienta", activo_id=h2.id, nombre_activo=h2.nombre, tipo="preventivo",
                                fecha_programada=datetime(2026, 1, 1), estado="realizado", coste_real=30),
        MovimientoMaterial(material_id=mat.id, tipo="salida", cantidad=12, obra_id=obra.id),
        MovimientoMaterial(material_id=mat.id, tipo="mostrador_salida", cantidad=8, obra_id=obra.id),
    ])
    db.commit()
    resp = client.post("/login", data={"username": "admin-gas", "password": "ClaveSegura123!"}, follow_redirects=False)
    client.cookies.set("mrd_token", resp.cookies["mrd_token"])
    client.cookies.set("mrd_csrf", generar_csrf_token())

    html = client.get("/informes").text
    assert 'id="informes-gasto"' in html and "Martillo caro" in html and "250.00 €" in html and "62 %" in html
    assert "Obra consumo" in html and "/informes/gasto-herramientas/excel" in html and "/informes/consumo-obras/excel" in html

    excel = client.get("/informes/gasto-herramientas/excel")
    assert excel.status_code == 200 and "spreadsheetml" in excel.headers["content-type"]
    ws = load_workbook(io.BytesIO(excel.content)).active
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    assert filas[0][0] == "GAS-1" and filas[0][5] == 250.0 and filas[1][0] == "GAS-2" and filas[1][4] == 30.0

    excel2 = client.get("/informes/consumo-obras/excel")
    ws2 = load_workbook(io.BytesIO(excel2.content)).active
    fila = list(ws2.iter_rows(min_row=2, values_only=True))[0]
    assert fila[0] == "2026-0099" and fila[3] == 20.0 and fila[4] == 1 and fila[5] == 1
